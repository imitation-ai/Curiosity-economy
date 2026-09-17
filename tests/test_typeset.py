import importlib.util
import tempfile
import unittest
from pathlib import Path
from docx import Document
from pypdf import PdfReader

spec = importlib.util.spec_from_file_location('typesetter', Path(__file__).resolve().parents[1] / 'scripts/typeset-book.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
module.register_fonts()


class TypesettingTests(unittest.TestCase):
    def test_rich_text_pagination_and_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'source.docx'
            output = Path(folder) / 'chapter.pdf'
            doc = Document()
            doc.add_heading('Chapter 1 — Curiosity & judgement', 0)
            p = doc.add_paragraph('Every question carries ')
            p.add_run('possibility').italic = True
            p.add_run(' — even when 2 < 3 & 3 > 2.')
            doc.add_heading('A different perspective', 1)
            for i in range(50):
                doc.add_paragraph(f'Paragraph {i}: “Keep the wording intact.” This is a complete paragraph that should move between pages without losing words or punctuation.')
            doc.save(source)
            result = module.typeset(source, output, 'Test book', 'Test Author', 1)
            self.assertGreater(result['pageCount'], 2)
            self.assertIn('<i>possibility</i>', result['blocks'][0]['html'])
            self.assertIn('&lt;', result['blocks'][0]['html'])
            self.assertEqual(result['blocks'][1]['type'], 'heading')
            pdf = PdfReader(output)
            self.assertEqual(tuple(pdf.pages[0].mediabox), (0, 0, 432, 648))
            self.assertIn('TEST AUTHOR', pdf.pages[2].extract_text())

    def test_new_tables_are_not_silently_discarded(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'source.docx'
            doc = Document()
            doc.add_paragraph('Chapter 1 — Title')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Do not lose this content.'
            doc.save(source)
            with self.assertRaisesRegex(ValueError, 'text-only'):
                module.typeset(source, Path(folder) / 'out.pdf', 'Test', 'Author', 0)


if __name__ == '__main__':
    unittest.main()
